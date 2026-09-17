from pathlib import Path
import json,math,hashlib
from openpyxl import load_workbook
from PIL import Image,ImageOps,ImageDraw
B=Path('F:/论文库/IGS/实验/过程/分类预测与SHAP_20260916');O=Path('F:/论文库/IGS/实验/结论/分类预测与SHAP_20260916');data=json.loads((B/'model_workbook.json').read_text(encoding='utf8'));w=load_workbook(O/'分类预测与SHAP分析结果.xlsx',data_only=True);checked=0
for name,t in data.items():
    sh=w[name];assert sh.max_row==len(t['rows'])+1
    for expected,actual in zip([t['columns']]+t['rows'],sh.iter_rows(values_only=True)):
        for e,v in zip(expected,actual):
            if isinstance(e,(int,float)):assert isinstance(v,(int,float)) and math.isclose(e,v,rel_tol=1e-8,abs_tol=1e-8),(name,e,v)
            else:assert (e or None)==(v or None),(name,e,v)
            checked+=1
check={'cells_verified':checked,'sheets':w.sheetnames,'raw_source_unchanged':hashlib.sha256(Path('F:/论文库/IGS/实验/过程/地质工程因素_20260916/pair_factors.pkl').read_bytes()).hexdigest()==json.loads((B/'audit.json').read_text(encoding='utf8'))['source_sha256']}
(B/'workbook_final_checks.json').write_text(json.dumps(check,ensure_ascii=False,indent=2),encoding='utf8')
# Contact sheets supplement direct figure/page inspection.
for pattern,out,width in [('表格预览_*.png','表格全表核查.png',580)]:
    paths=list(B.glob(pattern));canvas=Image.new('RGB',(width*2,300*4),'white');draw=ImageDraw.Draw(canvas)
    for i,path in enumerate(paths):
        im=Image.open(path);im.thumbnail((width-10,260));x=(i%2)*width;y=(i//2)*300;canvas.paste(im,(x,y+20));draw.text((x+5,y),str(i+1)+' '+path.stem,fill='black')
    canvas.save(B/out)
print(check)
